"""Coordinator contracts for Amazon review jobs."""

import tempfile
import unittest
from io import BytesIO
from pathlib import Path

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from engine.distributed.coordinator_models import Base, create_database_engine, create_session_factory
from engine.distributed.coordinator_server import create_coordinator_app
from engine.distributed.coordinator_store import CoordinatorStore
from engine.distributed.protocol import AgentLimits, hello_message, payload_checksum


class ReviewJobTests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        path = Path(self.directory.name) / "reviews.sqlite3"
        self.engine = create_database_engine(f"sqlite:///{path.as_posix()}")
        Base.metadata.create_all(self.engine)
        self.store = CoordinatorStore(create_session_factory(self.engine))

    def tearDown(self):
        self.engine.dispose()
        self.directory.cleanup()

    def test_review_job_requires_capable_agent_and_retains_result_and_samples(self):
        job = self.store.create_review_job({"source": "B012345678", "maxPages": 3})
        self.assertEqual(job["settings"]["channel"], "amazon_reviews")
        legacy = hello_message(client_id="legacy", display_name="Legacy", available_slots=1,
                               max_concurrent_inputs=1, limits=AgentLimits())
        legacy["capabilities"]["amazonReviews"] = False
        self.store.register_client(legacy)
        self.assertEqual(self.store.lease_tasks("legacy", 1), [])
        capable = hello_message(client_id="review-agent", display_name="Review agent", available_slots=1,
                                max_concurrent_inputs=1, limits=AgentLimits())
        self.store.register_client(capable)
        lease = self.store.lease_tasks("review-agent", 1)[0]
        review = {"reviewId": "R1", "body": "Clear design.", "rating": 5, "synthetic": False, "source": "amazon"}
        payload = {"jobId": job["id"], "taskId": lease["taskId"], "leaseId": lease["leaseId"],
                   "clientId": "review-agent", "products": [], "reviewData": {
                       "asin": "B012345678", "context": {"title": "Deer rug"}, "reviews": [review],
                       "reviewCount": 1, "pagesFetched": 2, "stopReason": "captcha"}}
        result = self.store.accept_result(lease["taskId"], "review-agent", lease["leaseId"],
                                          payload_checksum(payload), payload)
        self.assertEqual(result["status"], "accepted")
        sample = {"reviewId": "SYNTH-B012345678-001", "body": "The deer print is clear.",
                  "rating": 5, "synthetic": True, "source": "ai_sample", "verifiedPurchase": False}
        self.store.save_review_samples(job["id"], [sample])
        loaded = self.store.review_job(job["id"])
        self.assertEqual(loaded["status"], "completed")
        self.assertEqual(loaded["reviewData"]["reviews"], [review])
        self.assertEqual(loaded["samples"], [sample])

    def test_invalid_source_and_page_limit_are_rejected(self):
        with self.assertRaises(ValueError):
            self.store.create_review_job({"source": "https://example.com/dp/B012345678"})
        with self.assertRaises(ValueError):
            self.store.create_review_job({"source": "B012345678", "maxPages": 1001})

    def test_agent_failure_is_visible_on_review_job(self):
        job = self.store.create_review_job({"source": "B012345678"})
        self.store.register_client(hello_message(client_id="review-agent", display_name="Review agent",
                                                 available_slots=1, max_concurrent_inputs=1, limits=AgentLimits()))
        lease = self.store.lease_tasks("review-agent", 1)[0]
        self.store.fail_task("review-agent", {"taskId": lease["taskId"], "leaseId": lease["leaseId"],
                                              "error": {"code": "BROWSER_FAILED", "message": "Browser could not start", "retryable": False}})
        loaded = self.store.review_job(job["id"])
        self.assertEqual(loaded["status"], "partial")
        self.assertEqual(loaded["error"]["code"], "BROWSER_FAILED")

    def test_http_export_separates_real_and_synthetic_rows(self):
        database_path = Path(self.directory.name) / "review-api.sqlite3"
        app = create_coordinator_app(database_url=f"sqlite:///{database_path.as_posix()}")
        with TestClient(app) as client:
            created = client.post("/api/v1/review-jobs", json={"source": "B012345678", "maxPages": 1})
            self.assertEqual(created.status_code, 202)
            job_id = created.json()["id"]
            store = app.state.store
            store.register_client(hello_message(client_id="review-agent", display_name="Review agent",
                                                available_slots=1, max_concurrent_inputs=1, limits=AgentLimits()))
            lease = store.lease_tasks("review-agent", 1)[0]
            payload = {"jobId": job_id, "taskId": lease["taskId"], "leaseId": lease["leaseId"],
                       "clientId": "review-agent", "products": [], "reviewData": {
                           "asin": "B012345678", "reviews": [{"reviewId": "R1", "body": "Clear design.",
                           "rating": 5, "author": "Mia", "synthetic": False, "source": "amazon"}],
                           "context": {"title": "Deer rug"}}}
            store.accept_result(lease["taskId"], "review-agent", lease["leaseId"], payload_checksum(payload), payload)
            sample = {"reviewId": "SYNTH-B012345678-001", "body": "The deer print is clear.",
                      "rating": 5, "author": "Alex", "synthetic": True,
                      "source": "ai_sample", "verifiedPurchase": False, "promptVersion": "review_sample_v3"}
            saved = client.post(f"/api/v1/review-jobs/{job_id}/samples", json={"samples": [sample]})
            self.assertEqual(saved.status_code, 200)
            request = {"reviewIds": ["R1", sample["reviewId"]], "products": [{"id": "101", "handle": "deer-rug"}]}
            real = client.post(f"/api/v1/review-jobs/{job_id}/export", json={**request, "kind": "real"})
            preview = client.post(f"/api/v1/review-jobs/{job_id}/export", json={**request, "kind": "preview"})
            self.assertEqual(real.status_code, 200)
            self.assertEqual(preview.status_code, 200)
            real_book = load_workbook(BytesIO(real.content), read_only=True)
            preview_book = load_workbook(BytesIO(preview.content), read_only=True)
            self.assertEqual(real_book.active.max_row, 2)
            self.assertEqual(preview_book.active.max_row, 3)
            self.assertEqual(real_book.active["D2"].value, "R1")
            self.assertEqual(preview_book.active["M3"].value, "TRUE")


if __name__ == "__main__":
    unittest.main()
