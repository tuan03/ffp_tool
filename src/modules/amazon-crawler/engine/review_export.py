"""Three review XLSX layouts ported from the legacy Amazon Reviews tool."""

from __future__ import annotations

import random
import re
from datetime import datetime
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


JUDGEME_FIELDS = [
    "product_id", "shopify_product_id", "asin", "review_id", "review_date", "date",
    "rating", "reviewer_name", "name", "email", "body", "title", "verified",
    "variant", "picture_urls", "permalink", "source_url", "synthetic", "source", "import_note",
]
QA_SAMPLE_FIELDS = ["review_id", "review_date", "rating", "author", "body", "variant",
                    "synthetic", "quality_status", "quality_warnings", "prompt_version"]
QA_FIELDS = ["product_id", "product_handle", "reviewer_name", "reviewer_email", "title", "body",
             "rating", "review_date", "reply", "picture_urls", "review_id", "variant", "synthetic",
             "quality_status", "quality_warnings", "prompt_version"]


def _synthetic(review: dict[str, Any]) -> bool:
    return review.get("synthetic") is True or review.get("source") == "ai_sample"


def _date(review: dict[str, Any]) -> str:
    raw = str(review.get("review_date") or review.get("dateText") or "")
    if re.fullmatch(r"\d{4}-\d{2}-\d{2}", raw):
        return raw
    match = re.search(r"\b([A-Z][a-z]+ \d{1,2}, \d{4})\b", raw)
    if match:
        try:
            return datetime.strptime(match.group(1), "%B %d, %Y").date().isoformat()
        except ValueError:
            pass
    return ""


def _images(review: dict[str, Any]) -> list[str]:
    urls = []
    for image in review.get("images") or []:
        url = str(image.get("url") if isinstance(image, dict) else image).strip()
        if url.startswith(("https://", "http://")) and url not in urls:
            urls.append(url)
    return urls


def _prepare_reviews(reviews: list[dict[str, Any]], rng: random.Random, randomize: bool, minimum: int) -> list[dict[str, Any]]:
    if not randomize or not reviews:
        return list(reviews)
    count = rng.randint(max(1, min(minimum, len(reviews))), len(reviews))
    if count >= len(reviews):
        selected = list(reviews)
        rng.shuffle(selected)
        return selected
    target_average = rng.uniform(4.3, 4.7)
    valid = [review for review in reviews if isinstance(review.get("rating"), (int, float))]
    if len(valid) < count:
        return rng.sample(reviews, count)
    best: list[dict[str, Any]] | None = None
    best_score = float("inf")
    for _ in range(240):
        candidate = rng.sample(valid, count)
        average = sum(float(review["rating"]) for review in candidate) / count
        score = abs(average - target_average)
        if score < best_score:
            best, best_score = candidate, score
        if score <= 0.015:
            break
    selected = list(best or rng.sample(valid, count))
    rng.shuffle(selected)
    return selected


def _extra_images(reviews: list[dict[str, Any]], urls: list[str], rng: random.Random) -> dict[int, str]:
    eligible = [index for index, review in enumerate(reviews) if not _images(review)]
    count = min(len(eligible), len(urls))
    return dict(zip(rng.sample(eligible, count), rng.sample(urls, count)))


def _real_row(asin: str, review: dict[str, Any], product_id: str, source_url: str, extra: str) -> dict[str, Any]:
    author = str(review.get("author") or "Amazon Customer").strip()
    email_base = re.sub(r"[^a-z0-9]+", "", author.lower()) or re.sub(r"[^a-z0-9]+", "", str(review.get("reviewId") or "").lower())
    images = _images(review)
    if extra and extra not in images:
        images.append(extra)
    date = _date(review)
    return {
        "product_id": product_id, "shopify_product_id": product_id, "asin": asin,
        "review_id": review.get("reviewId") or "", "review_date": date, "date": date,
        "rating": int(round(float(review.get("rating") or 0))), "reviewer_name": author, "name": author,
        "email": f"{email_base}@gmail.com", "body": str(review.get("body") or "").strip(),
        "title": "", "verified": "TRUE" if review.get("verifiedPurchase") else "FALSE",
        "variant": str(review.get("variantText") or "").strip(), "picture_urls": ",".join(images),
        "permalink": review.get("permalink") or "", "source_url": source_url,
        "synthetic": "FALSE", "source": "amazon", "import_note": "",
    }


def _qa_row(review: dict[str, Any], product: dict[str, Any] | None, extra: str = "") -> dict[str, Any]:
    images = _images(review)
    if extra and extra not in images:
        images.append(extra)
    return {
        "product_id": str((product or {}).get("id") or ""), "product_handle": str((product or {}).get("handle") or ""),
        "reviewer_name": str(review.get("author") or ""), "reviewer_email": "",
        "title": str(review.get("title") or ""), "body": str(review.get("body") or ""),
        "rating": review.get("rating") or "", "review_date": _date(review), "reply": "",
        "picture_urls": ",".join(images), "review_id": review.get("reviewId") or "",
        "variant": review.get("variantText") or "", "synthetic": "TRUE" if _synthetic(review) else "FALSE",
        "quality_status": review.get("qualityStatus") or ("accepted" if _synthetic(review) else "source"),
        "quality_warnings": " | ".join(review.get("qualityWarnings") or []),
        "prompt_version": review.get("promptVersion") or "",
    }


def _safe_cell(value: Any) -> Any:
    if isinstance(value, str) and value.startswith("="):
        return f"'{value}"
    return value


def build_review_workbook(
    kind: str,
    asin: str,
    reviews: list[dict[str, Any]],
    products: list[dict[str, Any]],
    *,
    source_url: str = "",
    extra_picture_urls: list[str] | None = None,
    randomize_review_count: bool = False,
    min_reviews_per_product: int = 1,
    rng: random.Random | None = None,
) -> tuple[Workbook, int]:
    if kind not in {"real", "ai", "preview"}:
        raise ValueError("Unknown review export kind.")
    if not re.fullmatch(r"[A-Z0-9]{10}", asin):
        raise ValueError("A valid ASIN is required.")
    selected = [review for review in reviews if isinstance(review, dict) and str(review.get("body") or "").strip()]
    if kind == "real":
        selected = [review for review in selected if not _synthetic(review)]
    elif kind == "ai":
        selected = [review for review in selected if _synthetic(review)]
    if not selected:
        raise ValueError("No reviews of the selected kind are available.")
    chosen_rng = rng or random.Random()
    extra_urls = list(dict.fromkeys(url.strip() for url in extra_picture_urls or [] if url.startswith(("https://", "http://"))))
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = {"real": "Judge.me Reviews", "ai": "QA AI Samples", "preview": "QA Review Preview"}[kind]
    fields = JUDGEME_FIELDS if kind == "real" else (QA_SAMPLE_FIELDS if kind == "ai" and not products else QA_FIELDS)
    sheet.append(fields)
    row_count = 0
    targets: list[dict[str, Any] | None] = [products[0]] if kind == "ai" and products else [None] if not products else products
    for product in targets:
        product_reviews = _prepare_reviews(selected, chosen_rng, randomize_review_count if kind != "ai" else False, min_reviews_per_product)
        extras = _extra_images(product_reviews, extra_urls, chosen_rng)
        for index, review in enumerate(product_reviews):
            if kind == "real":
                row = _real_row(asin, review, str((product or {}).get("id") or ""), source_url, extras.get(index, ""))
                if row["rating"] < 1 or row["rating"] > 5:
                    continue
            else:
                row = _qa_row(review, product, extras.get(index, ""))
                if kind == "ai" and not products:
                    row["author"] = row["reviewer_name"]
            sheet.append([_safe_cell(row.get(field, "")) for field in fields])
            row_count += 1
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    fill = PatternFill("solid", fgColor="1F4E78" if kind == "real" else "5B21B6")
    for cell in sheet[1]:
        cell.fill = fill
        cell.font = Font(color="FFFFFF", bold=True)
    for index, field in enumerate(fields, 1):
        letter = get_column_letter(index)
        sheet.column_dimensions[letter].width = 72 if field == "body" else (42 if field in {"picture_urls", "quality_warnings"} else 20)
        if field in {"body", "picture_urls", "quality_warnings"}:
            for cell in sheet[letter]:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
    return workbook, row_count
